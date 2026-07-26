#!/usr/bin/env python3
"""Extract searchable text and file metadata from the labour invoice archive."""

from __future__ import annotations

import argparse
import hashlib
import html
import json
import mimetypes
import re
import subprocess
import tempfile
from email import policy
from email.parser import BytesParser
from pathlib import Path

from docx import Document
from pypdf import PdfReader


PDF_EXTENSIONS = {".pdf"}
IMAGE_EXTENSIONS = {".png", ".jpg", ".jpeg", ".heic"}
WORD_EXTENSIONS = {".docx"}
EMAIL_EXTENSIONS = {".eml"}


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def clean_text(value: str) -> str:
    value = html.unescape(value.replace("\x00", " "))
    value = re.sub(r"\r\n?", "\n", value)
    value = re.sub(r"[ \t]+", " ", value)
    value = re.sub(r"\n{3,}", "\n\n", value)
    return value.strip()


def source_date_from_name(name: str) -> str | None:
    match = re.match(r"^(\d{4}-\d{2}-\d{2})\s+-\s+", name)
    return match.group(1) if match else None


def ocr_image(path: Path) -> str:
    candidates: list[str] = []
    with tempfile.TemporaryDirectory(prefix="labour-invoice-image-ocr-") as temp_dir:
        for rotation in (0, 90, 180, 270):
            candidate_path = path
            if rotation:
                candidate_path = Path(temp_dir) / f"rotated-{rotation}.png"
                converted = subprocess.run(
                    ["sips", "-r", str(rotation), "-s", "format", "png", str(path), "--out", str(candidate_path)],
                    check=False,
                    capture_output=True,
                    text=True,
                )
                if converted.returncode != 0:
                    continue
            completed = subprocess.run(
                ["tesseract", str(candidate_path), "stdout", "--psm", "6"],
                check=False,
                capture_output=True,
                text=True,
            )
            extracted = clean_text(completed.stdout)
            if extracted:
                candidates.append(extracted)
    if not candidates:
        return ""
    return max(
        candidates,
        key=lambda value: (
            sum(keyword in value.lower() for keyword in ("invoice", "total", "paid", "receipt", "amount")),
            len(re.sub(r"\W", "", value)),
        ),
    )


def extract_pdf(path: Path) -> tuple[str, str, int | None]:
    try:
        reader = PdfReader(str(path))
        page_count = len(reader.pages)
        text = clean_text("\n\n".join((page.extract_text() or "") for page in reader.pages))
    except Exception as exc:
        return f"[PDF extraction failed: {exc}]", "pdf_error", None

    tokens = re.findall(r"\b\w+\b", text)
    single_character_ratio = (
        sum(1 for token in tokens if len(token) == 1) / len(tokens) if tokens else 0
    )
    if len(re.sub(r"\W", "", text)) >= 80 and single_character_ratio < 0.45:
        return text, "pdf_text", page_count

    ocr_pages: list[str] = []
    with tempfile.TemporaryDirectory(prefix="labour-invoice-ocr-") as temp_dir:
        prefix = Path(temp_dir) / "page"
        completed = subprocess.run(
            [
                "pdftoppm",
                "-f",
                "1",
                "-l",
                str(min(page_count, 4)),
                "-r",
                "220",
                "-png",
                str(path),
                str(prefix),
            ],
            check=False,
            capture_output=True,
            text=True,
        )
        if completed.returncode == 0:
            for rendered in sorted(Path(temp_dir).glob("page-*.png")):
                extracted = ocr_image(rendered)
                if extracted:
                    ocr_pages.append(extracted)
    ocr_text = clean_text("\n\n".join(ocr_pages))
    return (ocr_text or text), ("pdf_ocr" if ocr_text else "pdf_text_sparse"), page_count


def extract_docx(path: Path) -> str:
    document = Document(str(path))
    blocks: list[str] = [paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()]
    for table in document.tables:
        for row in table.rows:
            cells = [clean_text(cell.text) for cell in row.cells]
            if any(cells):
                blocks.append(" | ".join(cells))
    return clean_text("\n".join(blocks))


def extract_email(path: Path) -> str:
    message = BytesParser(policy=policy.default).parsebytes(path.read_bytes())
    blocks = [
        f"Subject: {message.get('subject', '')}",
        f"From: {message.get('from', '')}",
        f"Date: {message.get('date', '')}",
    ]
    for part in message.walk():
        if part.get_content_maintype() == "multipart":
            continue
        if part.get_content_type() not in {"text/plain", "text/html"}:
            continue
        try:
            content = part.get_content()
        except Exception:
            continue
        if part.get_content_type() == "text/html":
            content = re.sub(r"<[^>]+>", " ", content)
        blocks.append(content)
    return clean_text("\n".join(blocks))


def extract_image(path: Path) -> tuple[str, str]:
    if path.suffix.lower() != ".heic":
        return ocr_image(path), "image_ocr"
    with tempfile.TemporaryDirectory(prefix="labour-invoice-heic-") as temp_dir:
        converted = Path(temp_dir) / f"{path.stem}.png"
        completed = subprocess.run(
            ["sips", "-s", "format", "png", str(path), "--out", str(converted)],
            check=False,
            capture_output=True,
            text=True,
        )
        if completed.returncode != 0 or not converted.exists():
            return "", "heic_conversion_failed"
        return ocr_image(converted), "image_ocr"


def extract_file(path: Path) -> dict:
    suffix = path.suffix.lower()
    text = ""
    extraction_method = "unsupported"
    page_count = None
    error = None
    try:
        if suffix in PDF_EXTENSIONS:
            text, extraction_method, page_count = extract_pdf(path)
        elif suffix in WORD_EXTENSIONS:
            text, extraction_method = extract_docx(path), "docx_text"
        elif suffix in EMAIL_EXTENSIONS:
            text, extraction_method = extract_email(path), "eml_text"
        elif suffix in IMAGE_EXTENSIONS:
            text, extraction_method = extract_image(path)
        elif suffix == ".xlsx":
            extraction_method = "xlsx_pending_artifact_tool"
    except Exception as exc:
        error = f"{type(exc).__name__}: {exc}"

    return {
        "source_path": str(path),
        "source_filename": path.name,
        "source_date": source_date_from_name(path.name),
        "extension": suffix.lstrip("."),
        "mime_type": mimetypes.guess_type(path.name)[0] or "application/octet-stream",
        "file_size": path.stat().st_size,
        "sha256": sha256_file(path),
        "page_count": page_count,
        "extraction_method": extraction_method,
        "text_length": len(text),
        "text": text,
        "error": error,
    }


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("input_dir", type=Path)
    parser.add_argument("output_json", type=Path)
    args = parser.parse_args()

    files = sorted(path for path in args.input_dir.iterdir() if path.is_file())
    rows = [extract_file(path) for path in files]
    args.output_json.parent.mkdir(parents=True, exist_ok=True)
    args.output_json.write_text(json.dumps(rows, ensure_ascii=False, indent=2), encoding="utf-8")

    methods: dict[str, int] = {}
    for row in rows:
        methods[row["extraction_method"]] = methods.get(row["extraction_method"], 0) + 1
    print(json.dumps({"files": len(rows), "methods": methods}, indent=2))


if __name__ == "__main__":
    main()
